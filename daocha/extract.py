"""Extract 【n】 markers from PDF or Word body text."""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

# Canonical slot: 【1】 【2】 …  (ASCII digits; fullwidth digits are normalized)
VALID_INNER_RE = re.compile(r"^\d+$")
ANY_MARKER_RE = re.compile(r"【[^】]*】")
SPACED_NUMBER_RE = re.compile(r"【\s*(\d+)\s*】")
RANGE_INNER_RE = re.compile(r"[-–—,/]|到|至")
FULLWIDTH_DIGITS = str.maketrans("０１２３４５６７８９", "0123456789")

# Re-exported name used by manuscript.py
MARKER_RE = ANY_MARKER_RE


class MarkerExtractError(ValueError):
    def __init__(self, report: dict[str, Any]):
        self.report = report
        super().__init__(format_extract_report(report))


def normalize_manuscript_text(text: str) -> str:
    """Unify digits and close up 【 12 】 → 【12】."""
    text = text.translate(FULLWIDTH_DIGITS)
    return SPACED_NUMBER_RE.sub(lambda m: f"【{m.group(1)}】", text)


def read_manuscript_text(path: Path | str) -> str:
    return "\n".join(chunk for _loc, chunk in iter_text_chunks(path))


def iter_text_chunks(path: Path | str) -> list[tuple[str, str]]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".doc":
        raise ValueError("不支持旧版 .doc，请另存为 .docx 或 PDF。")
    if suffix == ".pdf":
        import fitz

        doc = fitz.open(source)
        chunks: list[tuple[str, str]] = []
        for index, page in enumerate(doc, start=1):
            chunks.append((f"第 {index} 页", page.get_text("text") or ""))
        doc.close()
        return chunks
    if suffix == ".docx":
        from docx import Document

        document = Document(str(source))
        chunks = []
        for index, para in enumerate(document.paragraphs, start=1):
            if para.text.strip():
                chunks.append((f"正文第 {index} 段", para.text))
        for t_i, table in enumerate(document.tables, start=1):
            for r_i, row in enumerate(table.rows, start=1):
                cell_text = " ".join(cell.text for cell in row.cells if cell.text.strip())
                if cell_text:
                    chunks.append((f"表格 {t_i} 第 {r_i} 行", cell_text))
        return chunks
    if suffix in {".txt", ".md"}:
        return [("全文", source.read_text(encoding="utf-8"))]
    raise ValueError(f"请使用 PDF、.docx、.txt 或 .md（当前是 {suffix or source.name}）")


def _snippet(text: str, start: int, end: int, radius: int = 40) -> str:
    a = max(0, start - radius)
    b = min(len(text), end + radius)
    return re.sub(r"\s+", " ", text[a:b]).strip()


def analyze_markers(text: str, *, location: str = "") -> dict[str, Any]:
    text = normalize_manuscript_text(text)
    issues: list[dict[str, Any]] = []
    valid: list[dict[str, str]] = []

    for match in ANY_MARKER_RE.finditer(text):
        raw = match.group(0)
        inner = raw[1:-1]
        loc = location
        snippet = _snippet(text, match.start(), match.end())
        if not inner.strip() or "\ufffd" in inner:
            issues.append(
                {"type": "empty", "marker": raw, "location": loc, "context": snippet}
            )
            continue
        if RANGE_INNER_RE.search(inner) or inner.count("【"):
            issues.append(
                {"type": "range", "marker": raw, "location": loc, "context": snippet}
            )
            continue
        if not VALID_INNER_RE.fullmatch(inner.strip()):
            issues.append(
                {"type": "invalid", "marker": raw, "location": loc, "context": snippet}
            )
            continue
        marker = f"【{inner.strip()}】"
        valid.append(
            {
                "marker": marker,
                "context": snippet,
                "location": loc,
            }
        )

    by_marker: dict[str, list[dict[str, str]]] = defaultdict(list)
    for item in valid:
        by_marker[item["marker"]].append(item)
    unique: list[dict[str, str]] = []
    seen: set[str] = set()
    for item in valid:
        marker = item["marker"]
        if marker in seen:
            continue
        seen.add(marker)
        copies = by_marker[marker]
        if len(copies) > 1:
            issues.append(
                {
                    "type": "duplicate",
                    "marker": marker,
                    "count": len(copies),
                    "location": "；".join(
                        sorted({c.get("location") or "" for c in copies if c.get("location")})
                    ),
                    "context": copies[0].get("context") or "",
                }
            )
            continue
        unique.append(item)

    return {
        "markers": unique,
        "issues": issues,
        "valid_slots": len(unique),
        "raw_hits": len(valid),
    }


def analyze_file(path: Path | str) -> dict[str, Any]:
    merged_issues: list[dict[str, Any]] = []
    merged_valid: list[dict[str, str]] = []
    for location, chunk in iter_text_chunks(path):
        part = analyze_markers(chunk, location=location)
        merged_issues.extend(part["issues"])
        merged_valid.extend(part["markers"])

    by_marker: dict[str, list[dict[str, str]]] = defaultdict(list)
    for item in merged_valid:
        by_marker[item["marker"]].append(item)
    unique: list[dict[str, str]] = []
    seen: set[str] = set()
    extra_dupes: list[dict[str, Any]] = []
    for item in merged_valid:
        marker = item["marker"]
        if marker in seen:
            continue
        seen.add(marker)
        copies = by_marker[marker]
        if len(copies) > 1:
            extra_dupes.append(
                {
                    "type": "duplicate",
                    "marker": marker,
                    "count": len(copies),
                    "location": "；".join(
                        c.get("location") or "" for c in copies if c.get("location")
                    ),
                    "context": copies[0].get("context") or "",
                }
            )
            continue
        unique.append(item)

    issues = [i for i in merged_issues if i.get("type") != "duplicate"] + extra_dupes
    return {
        "markers": unique,
        "issues": issues,
        "valid_slots": len(unique),
        "raw_hits": len(merged_valid),
    }


def format_extract_report(report: dict[str, Any]) -> str:
    lines = [
        "【】抽取未通过。请改成一个坑一个号：【1】【2】【3】，不要写【1-3】，不要重复用同一个号。",
        f"合法引文位：{report.get('valid_slots', 0)}",
        f"问题：{len(report.get('issues') or [])} 处",
        "",
    ]
    labels = {
        "empty": "空括号（数字丢失，常见于 Word 转 PDF）",
        "range": "合并/范围编号",
        "duplicate": "同一个号用了两次",
        "invalid": "不是纯数字编号",
    }
    for issue in report.get("issues") or []:
        kind = labels.get(issue.get("type"), issue.get("type"))
        loc = issue.get("location") or ""
        marker = issue.get("marker") or ""
        ctx = issue.get("context") or ""
        extra = f" ×{issue['count']}" if issue.get("count") else ""
        lines.append(f"- {kind}{extra}: {marker}  {loc}")
        if ctx:
            lines.append(f"  附近：{ctx}")
    return "\n".join(lines).strip()


def extract_markers(path: Path | str) -> list[dict[str, str]]:
    report = analyze_file(path)
    if report["issues"]:
        raise MarkerExtractError(report)
    if not report["markers"]:
        raise MarkerExtractError(
            {
                **report,
                "issues": [
                    {
                        "type": "invalid",
                        "marker": "",
                        "location": "",
                        "context": "没有找到【1】【2】这种标记",
                    }
                ],
            }
        )
    return report["markers"]


def empty_citation(index: int, marker: str, context: str, location: str = "") -> dict[str, Any]:
    return {
        "id": f"C{index:03d}",
        "markerIndex": index,
        "marker": marker,
        "section": "",
        "location": location,
        "claim": "",
        "recommended": "",
        "title": "",
        "journal": "",
        "year": None,
        "doi": "",
        "url": "",
        "excerpt": context[:400],
        "sourceParagraph": "",
        "strength": "",
        "status": "pending",
        "userDecision": "",
        "userNote": "",
    }


def build_registry(
    markers: list[dict[str, str]],
    *,
    project_id: str,
    manuscript_name: str,
    title: str = "",
    target_journal: str = "Journal of Environmental Management",
    max_same_source: int = 2,
    ban_mdpi: bool = True,
) -> dict[str, Any]:
    citations = [
        empty_citation(
            i,
            item["marker"],
            item.get("context") or "",
            item.get("location") or "",
        )
        for i, item in enumerate(markers, start=1)
    ]
    return {
        "meta": {
            "projectId": project_id,
            "manuscriptTitle": title,
            "manuscriptFile": manuscript_name,
            "targetJournal": target_journal,
            "workflowPhase": 1,
            "phaseLabel": "核查中",
            "lastUpdated": date.today().isoformat(),
            "constraints": {
                "banMdpi": True,
                "maxSameSource": max_same_source,
            },
            "pdfMarkerCount": len(citations),
            "registryEntryCount": len(citations),
            "oneToOneMarkers": True,
        },
        "citations": citations,
    }
